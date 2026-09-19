"""Sinh toàn bộ tệp trong tests/fixtures/ một cách tất định.

Chạy:  cd tests/fixtures && pip install -r requirements.txt && python generate.py

Kèm --with-oversize để sinh thêm file .xlsx > 10 MB (không commit, xem .gitignore).

Script này vừa sinh dữ liệu vừa TỰ KIỂM ĐỊNH: mọi con số trong expected.json đều
do qt_reference.py tính ra, rồi được đối chiếu với các assert viết tay bên dưới
(mục KIEM_DINH). Nếu hai bên lệch nhau, script dừng - không xuất file sai.
"""

from __future__ import annotations

import argparse
import json
import os
import random
import shutil
import sys
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook

import dataset as ds
from qt_reference import (
    Member,
    Period,
    anniversary,
    eligible_list,
    eligible_milestone,
    gap_ranges,
    milestones,
    missed_milestone,
    next_milestone,
    overlaps,
    party_age,
    period_status,
    upcoming_period,
    vietnamese_sort_key,
)

if hasattr(sys.stdout, "reconfigure"):  # console Windows mặc định cp1252
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent
EXCEL = ROOT / "excel"
DATA = ROOT / "data"
SQL = ROOT / "sql"

DATE_FMT = "DD/MM/YYYY"

# Tên bảng / cột dùng cho seed SQL. Nếu migration của Backend (T06) dùng tên khác,
# sửa ở ĐÚNG MỘT chỗ này rồi chạy lại generate.py.
DB = {
    "member_table": "party_members",
    "member_cols": ("id", "full_name", "date_of_birth", "gender",
                    "official_admission_date", "created_at"),
    "period_table": "award_periods",
    "period_cols": ("id", "name", "from_day", "from_month", "to_day", "to_month",
                    "created_at"),
    "setting_table": "app_settings",
    "setting_cols": ("id", "start_years", "end_years", "step_years", "unit_name",
                     "created_at"),
}
SEED_CREATED_AT = "2026-09-19T00:00:00+07:00"


# ---------------------------------------------------------------- tiện ích


def iso(d: date | None) -> str | None:
    return d.isoformat() if d else None


def vn(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def guid(prefix: str, index: int) -> str:
    return f"{prefix}0000000-0000-4000-8000-{index:012d}"


def sql_str(v: str | None) -> str:
    if v is None:
        return "NULL"
    return "'" + v.replace("'", "''") + "'"


def sql_date(d: date | None) -> str:
    return f"'{d.isoformat()}'" if d else "NULL"


# ---------------------------------------------------------------- tính kịch bản


def scenario(members: list[Member], periods: list[Period], settings: dict,
             today: date, years: list[int]) -> dict:
    ms = milestones(settings["start_years"], settings["end_years"], settings["step_years"])

    up = upcoming_period(periods, today)
    upcoming = None
    if up:
        p, y = up
        f, t = p.bind(y)
        status, days = period_status(p, today)
        upcoming = {
            "code": p.code, "name": p.name, "year": y,
            "boundFrom": iso(f), "boundTo": iso(t),
            "boundFromDisplay": vn(f), "boundToDisplay": vn(t),
            # QT11 chỉ nói về năm hiện tại; đợt của năm sau luôn là "Sắp tới"
            "status": status if y == today.year else "Sắp tới",
            "daysLeft": days if y == today.year else (p.bind(y)[0] - today).days,
        }

    period_statuses = []
    for p in sorted(periods, key=lambda x: x.sort_key):
        status, days = period_status(p, today)
        period_statuses.append({"code": p.code, "name": p.name,
                                "status": status, "daysLeft": days})

    eligible: dict[str, dict] = {}
    for p in sorted(periods, key=lambda x: x.sort_key):
        per_year = {}
        for y in years:
            rows = eligible_list(members, p, y, ms)
            by_ms: dict[str, int] = {}
            for _, n, _a in rows:
                by_ms[str(n)] = by_ms.get(str(n), 0) + 1
            f, t = p.bind(y)
            per_year[str(y)] = {
                "boundFrom": iso(f), "boundTo": iso(t),
                "total": len(rows),
                "byMilestone": by_ms,
                "rows": [{"code": m.code, "fullName": m.full_name,
                          "milestone": n, "anniversary": iso(a),
                          "anniversaryDisplay": vn(a)} for m, n, a in rows],
            }
        eligible[p.code] = {"name": p.name, "byYear": per_year}

    missed: dict[str, dict] = {}
    for y in years:
        rows = []
        for m in members:
            hit = missed_milestone(m.admission, periods, y, ms)
            if hit:
                n, a, label = hit
                rows.append({"code": m.code, "fullName": m.full_name,
                             "milestone": n, "anniversary": iso(a),
                             "anniversaryDisplay": vn(a), "gapLabel": label})
        rows.sort(key=lambda r: (r["milestone"], vietnamese_sort_key(r["fullName"])))
        missed[str(y)] = {"total": len(rows), "rows": rows}

    gaps = {}
    for y in years:
        gaps[str(y)] = [{"from": iso(f), "to": iso(t),
                         "fromDisplay": vn(f), "toDisplay": vn(t), "label": lbl}
                        for f, t, lbl in gap_ranges(periods, y)]

    return {
        "today": iso(today),
        "settings": settings,
        "milestones": ms,
        "milestoneCount": len(ms),
        "upcomingPeriod": upcoming,
        "periodStatuses": period_statuses,
        "eligibleByPeriod": eligible,
        "missedByYear": missed,
        "gapsByYear": gaps,
        "badgeCurrentYear": missed[str(today.year)]["total"],
        "overlapWarnings": overlaps(periods, today.year),
    }


def member_display_rows(members: list[Member], settings: dict, today: date) -> list[dict]:
    ms = milestones(settings["start_years"], settings["end_years"], settings["step_years"])
    out = []
    for m in members:
        age = party_age(m.admission, today)
        nm = next_milestone(m.admission, today, ms)
        na = anniversary(m.admission, nm) if nm else None
        out.append({
            "code": m.code,
            "fullName": m.full_name,
            "dateOfBirth": iso(m.date_of_birth),
            "dateOfBirthDisplay": vn(m.date_of_birth) or "—",
            "gender": m.gender,
            "genderDisplay": m.gender or "—",
            "officialAdmissionDate": iso(m.admission),
            "officialAdmissionDateDisplay": vn(m.admission),
            "partyAge": age,
            "nextMilestone": nm,
            "nextMilestoneDisplay": str(nm) if nm else "—",
            "nextAnniversary": iso(na),
            "nextAnniversaryDisplay": vn(na) or "—",
            "intent": m.intent,
        })
    return out


# ---------------------------------------------------------------- KIEM_DINH
# Các assert dưới đây là con số QC suy ra bằng tay từ tài liệu nghiệp vụ.
# Chúng canh cho qt_reference.py không tự do trôi.


def self_check(core: list[Member], scen: dict) -> None:
    by_code = {m.code: m for m in core}
    md = milestones(30, 90, 5)
    m10 = milestones(30, 90, 10)

    # QT1
    assert md == [30, 35, 40, 45, 50, 55, 60, 65, 70, 75, 80, 85, 90], md
    assert len(md) == 13
    assert m10 == [30, 40, 50, 60, 70, 80, 90], m10
    assert milestones(30, 90, 61) == [30], "Bước lớn hơn khoảng -> chỉ còn mốc đầu"
    assert milestones(30, 30, 5) == [30], "Bắt đầu == Kết thúc -> đúng 1 mốc"

    # QT2 - 29/02
    assert anniversary(date(1996, 2, 29), 30) == date(2026, 2, 28), "năm đích không nhuận"
    assert anniversary(date(1988, 2, 29), 40) == date(2028, 2, 29), "năm đích nhuận"
    assert anniversary(date(1996, 10, 1), 30) == date(2026, 10, 1)

    # QT3 / QT3a
    assert party_age(by_code["M01"].admission, ds.T0) == 91
    assert next_milestone(by_code["M01"].admission, ds.T0, md) is None
    assert party_age(by_code["M02"].admission, ds.T0) == 90
    assert next_milestone(by_code["M02"].admission, ds.T0, md) is None
    assert party_age(by_code["V01"].admission, ds.T0) == 0
    assert next_milestone(by_code["V01"].admission, ds.T0, md) == 30
    assert party_age(by_code["N01"].admission, ds.T0) == 28
    assert next_milestone(by_code["N01"].admission, ds.T0, md) == 30
    # tuổi đảng ngay trước và ngay sau ngày kỷ niệm
    assert party_age(date(1996, 9, 20), ds.T0) == 29, "chưa tới kỷ niệm thứ 30"
    assert party_age(date(1996, 9, 19), ds.T0) == 30, "đúng ngày kỷ niệm đã tính"
    assert party_age(date(1996, 2, 29), date(2026, 2, 27)) == 29
    assert party_age(date(1996, 2, 29), date(2026, 2, 28)) == 30, "29/02 thu về 28/02"

    # QT4 - biên đợt
    p4 = next(p for p in ds.PERIODS_MAIN if p.code == "P4")
    p1 = next(p for p in ds.PERIODS_MAIN if p.code == "P1")
    assert eligible_milestone(by_code["B01"].admission, p4, 2026, md) == 30
    assert eligible_milestone(by_code["B02"].admission, p4, 2026, md) == 30
    assert eligible_milestone(by_code["B03"].admission, p4, 2026, md) is None
    assert eligible_milestone(by_code["B04"].admission, p4, 2026, md) is None
    assert eligible_milestone(by_code["B05"].admission, p1, 2026, md) == 30
    assert eligible_milestone(by_code["B06"].admission, p1, 2026, md) == 30
    assert eligible_milestone(by_code["L01"].admission, p1, 2026, md) == 30
    assert eligible_milestone(by_code["M02"].admission,
                              next(p for p in ds.PERIODS_MAIN if p.code == "P2"),
                              2026, md) == 90

    e = scen["eligibleByPeriod"]
    assert e["P4"]["byYear"]["2026"]["total"] == 6, e["P4"]["byYear"]["2026"]["total"]
    assert e["P4"]["byYear"]["2026"]["byMilestone"] == {"30": 3, "35": 1, "40": 1, "45": 1}
    # UC-11: sắp theo Mốc rồi Họ tên (thứ tự chữ cái tiếng Việt: Đ < N < T)
    names30 = [r["fullName"] for r in e["P4"]["byYear"]["2026"]["rows"] if r["milestone"] == 30]
    assert names30 == ["Đào Văn Ân", "Nguyễn Văn An", "Trần Thị Bình"], names30
    assert e["P4"]["byYear"]["2027"]["total"] >= 1
    assert any(r["code"] == "N01" for r in e["P4"]["byYear"]["2027"]["rows"]), \
        "N01 chỉ đủ điều kiện ở năm sau"

    # L02: không có mốc nào rơi vào 2026 -> không đủ điều kiện và cũng không bị sót
    for code in e:
        assert not any(r["code"] == "L02" for r in e[code]["byYear"]["2026"]["rows"])
    assert not any(r["code"] == "L02" for r in scen["missedByYear"]["2026"]["rows"])
    assert any(r["code"] == "L02" for r in e["P1"]["byYear"]["2028"]["rows"]), \
        "L02 đủ điều kiện Đợt 3/2 năm 2028 với mốc 40 ngày 29/02/2028"

    # QT7 - nhãn khoảng trống
    labels = {r["code"]: r["gapLabel"] for r in scen["missedByYear"]["2026"]["rows"]}
    assert labels["B07"] == "Trước đợt đầu tiên", labels
    assert labels["B08"] == "Giữa Đợt 3/2 và Đợt 19/5", labels
    assert labels["G05"] == "Giữa Đợt 19/5 và Đợt 2/9", labels
    assert labels["M03"] == "Giữa Đợt 19/5 và Đợt 2/9", labels
    assert labels["B03"] == "Giữa Đợt 2/9 và Đợt 7/11", labels
    assert labels["B04"] == "Sau đợt cuối cùng", labels
    assert labels["S04"] == "Giữa Đợt 19/5 và Đợt 2/9", labels
    assert scen["missedByYear"]["2026"]["total"] == 7, scen["missedByYear"]["2026"]["total"]
    assert scen["badgeCurrentYear"] == 7

    # QT8 / QT11 tại T0
    up = scen["upcomingPeriod"]
    assert up["code"] == "P4" and up["year"] == 2026, up
    assert up["status"] == "Sắp tới" and up["daysLeft"] == 12, up
    st = {s["code"]: (s["status"], s["daysLeft"]) for s in scen["periodStatuses"]}
    assert st["P1"] == ("Đã qua", None) and st["P3"] == ("Đã qua", None), st

    # QT6 - cảnh báo
    assert scen["overlapWarnings"] == [], "bộ 4 đợt chính không chồng lấn"
    assert len(scen["gapsByYear"]["2026"]) == 5, scen["gapsByYear"]["2026"]
    assert overlaps(ds.PERIODS_OVERLAP, 2026) == [("Đợt A", "Đợt B")]
    assert gap_ranges(ds.PERIODS_FULL_COVER, 2026) == [], "bộ phủ kín không có khoảng trống"

    # Biên đợt rơi đúng 29/02 (QT4 phía đợt)
    pl = ds.PERIODS_LEAP_EDGE[0]
    assert pl.bind(2026) == (date(2026, 2, 28), date(2026, 3, 5)), pl.bind(2026)
    assert pl.bind(2028) == (date(2028, 2, 29), date(2028, 3, 5)), pl.bind(2028)
    assert eligible_milestone(by_code["L01"].admission, pl, 2026, md) == 30, \
        "28/02/2026 đúng bằng Từ ngày đã thu về"

    print("  self_check: OK")


# ---------------------------------------------------------------- xuất Excel


def write_xlsx(path: Path, rows: list[tuple], header: list[str] | None,
               date_cols: set[int] | None = None) -> None:
    """Ghi 1 sheet. date_cols = chỉ số cột (0-based) cần ghi ở KIỂU NGÀY thật của Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DanhSach"
    if header:
        ws.append(header)
    for r in rows:
        ws.append(list(r))
    if date_cols:
        for row in ws.iter_rows(min_row=2 if header else 1):
            for idx in date_cols:
                if idx < len(row) and isinstance(row[idx].value, date):
                    row[idx].number_format = DATE_FMT
    for col, width in zip("ABCD", (28, 14, 12, 26)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def member_row_text(m: Member) -> tuple:
    return (m.full_name, vn(m.date_of_birth), m.gender or "", vn(m.admission))


def member_row_date(m: Member) -> tuple:
    return (m.full_name, m.date_of_birth, m.gender or "", m.admission)


def build_excel(core: list[Member], bulk: list[Member]) -> list[dict]:
    EXCEL.mkdir(parents=True, exist_ok=True)
    manifest: list[dict] = []

    def rec(name: str, valid: int, errors: int, purpose: str, note: str = "") -> None:
        manifest.append({"file": name, "validRows": valid, "errorRows": errors,
                         "purpose": purpose, "note": note})

    # UC-25 file mẫu
    write_xlsx(EXCEL / "mau-dang-vien.xlsx",
               [("Nguyễn Văn Mẫu", "12/03/1974", "Nam", "01/10/1996"),
                ("Trần Thị Mẫu", "05/07/1973", "Nữ", "07/11/1996")],
               ds.HEADER)
    rec("mau-dang-vien.xlsx", 2, 0,
        "UC-25 - đối chiếu file mẫu do hệ thống tạo ra (đúng 4 cột, đúng thứ tự, ngày dd/MM/yyyy)")

    # Bộ lõi - ngày ở dạng chuỗi
    write_xlsx(EXCEL / "core-hop-le.xlsx", [member_row_text(m) for m in core], ds.HEADER)
    rec("core-hop-le.xlsx", len(core), 0,
        "Bộ lõi 31 ca biên, ngày ghi dạng CHUỖI dd/MM/yyyy. File chính cho E2E-1.")

    # Bộ lõi - ngày ở kiểu ngày thật của Excel
    write_xlsx(EXCEL / "core-hop-le-ngay-kieu-date.xlsx",
               [member_row_date(m) for m in core], ds.HEADER, date_cols={1, 3})
    rec("core-hop-le-ngay-kieu-date.xlsx", len(core), 0,
        "Cùng dữ liệu nhưng ô ngày là KIỂU NGÀY của Excel (số serial), không phải chuỗi. "
        "Trình đọc phải nhận cả hai kiểu.")

    # Bộ lớn
    write_xlsx(EXCEL / "bulk-1200.xlsx", [member_row_text(m) for m in bulk], ds.HEADER)
    rec("bulk-1200.xlsx", len(bulk), 0,
        "Phân trang / tìm kiếm / hiệu năng. Ngày chính thức 2015-2020 nên KHÔNG ảnh hưởng "
        "con số đủ điều kiện của bộ lõi ở các năm 2025-2030.")

    # E2E-2: đúng 4 dòng lỗi
    four = [("E-HOTEN-TRONG",), ("E-NGAYCT-TRONG",), ("E-NGAYCT-SAIDANG",),
            ("E-NGAYCT-TUONGLAI",)]
    codes = {c[0] for c in four}
    err4 = [r for r in ds.ERR_ROWS if r[4] in codes]
    assert len(err4) == 4
    rows = list(ds.OK_ROWS) + [(r[0], r[1], r[2], r[3]) for r in err4]
    write_xlsx(EXCEL / "loi-4-dong.xlsx", rows, ds.HEADER)
    rec("loi-4-dong.xlsx", len(ds.OK_ROWS), 4,
        "E2E-2 - xem trước phải báo 'Sẽ thêm 6 người mới · 4 dòng lỗi bị bỏ qua'. "
        "Dòng lỗi ở các dòng Excel 8, 9, 10, 11.")

    # Mỗi loại lỗi một dòng
    rows = [(r[0], r[1], r[2], r[3]) for r in ds.ERR_ROWS]
    rows += [ds.OK_ROWS[0], ds.OK_ROW_TODAY]
    write_xlsx(EXCEL / "loi-moi-loai-mot-dong.xlsx", rows, ds.HEADER)
    rec("loi-moi-loai-mot-dong.xlsx", 2, len(ds.ERR_ROWS),
        "QT9 - mỗi loại lỗi cấp dòng một dòng, kèm 2 dòng hợp lệ (dòng cuối là biên "
        "'Ngày chính thức đúng bằng hôm nay'). Bảng lỗi phải nêu đúng số dòng Excel và lý do.")

    # Chuẩn hoá - hành vi chưa chốt
    rows = [(r[0], r[1], r[2], r[3]) for r in ds.NORMALIZE_ROWS]
    write_xlsx(EXCEL / "bien-chuan-hoa.xlsx", rows, ds.HEADER)
    rec("bien-chuan-hoa.xlsx", 0, 0,
        "Các dòng mà tài liệu CHƯA quy định (OQ-4 trim họ tên, OQ-5 hoa/thường giới tính, "
        "OQ-6 ngày d/M/yyyy). Kết quả mong đợi chỉ chốt sau khi CEO trả lời.")

    # Lỗi cấp file
    write_xlsx(EXCEL / "loi-sai-cot.xlsx",
               [("01/10/1996", "Nguyễn Sai Cột", "Nam", "12/03/1974", "ghi chú")],
               ["Ngày vào Đảng chính thức", "Họ tên", "Giới tính", "Ngày sinh", "Ghi chú"])
    rec("loi-sai-cot.xlsx", 0, 0,
        "QT9 lỗi cấp file - 5 cột, sai thứ tự. Phải bị chặn ngay ở bước 1, không sang xem trước.")

    wb = Workbook()
    wb.active.title = "DanhSach"
    wb.save(EXCEL / "loi-rong.xlsx")
    rec("loi-rong.xlsx", 0, 0, "QT9 lỗi cấp file - sheet hoàn toàn rỗng, không có cả tiêu đề.")

    write_xlsx(EXCEL / "loi-chi-co-tieu-de.xlsx", [], ds.HEADER)
    rec("loi-chi-co-tieu-de.xlsx", 0, 0,
        "QT9 lỗi cấp file - đúng tiêu đề nhưng 0 dòng dữ liệu. Thông báo phải khác "
        "thông báo của file rỗng hoàn toàn (OQ-7).")

    (EXCEL / "loi-khong-phai-xlsx.xlsx").write_text(
        "Đây là tệp văn bản thuần, chỉ đổi phần mở rộng thành .xlsx.\n"
        "Trình đọc phải báo lỗi định dạng, không được văng lỗi 500.\n",
        encoding="utf-8")
    rec("loi-khong-phai-xlsx.xlsx", 0, 0,
        "QT9 lỗi cấp file - không phải tệp zip/xlsx. Phải trả 400 với thông báo tiếng Việt, "
        "không phải 500.")

    (EXCEL / "loi-dinh-dang-csv.csv").write_text(
        "Họ tên,Ngày sinh,Giới tính,Ngày vào Đảng chính thức\n"
        "Nguyễn Văn Csv,12/03/1974,Nam,01/10/1996\n",
        encoding="utf-8-sig")
    rec("loi-dinh-dang-csv.csv", 0, 0,
        "QT9 lỗi cấp file - .csv phải bị từ chối vì chỉ nhận .xlsx.")

    return manifest


def build_oversize() -> None:
    """Sinh file .xlsx hợp lệ nhưng > 10 MB bằng cách nhồi một phần nhị phân khó nén."""
    src = EXCEL / "core-hop-le.xlsx"
    dst = EXCEL / "loi-qua-10mb.xlsx"
    shutil.copyfile(src, dst)
    rnd = random.Random(20260919)
    pad = rnd.randbytes(11 * 1024 * 1024)
    with zipfile.ZipFile(dst, "a", compression=zipfile.ZIP_STORED) as z:
        z.writestr("xl/media/padding.bin", pad)
    size_mb = dst.stat().st_size / 1024 / 1024
    assert size_mb > 10, size_mb
    print(f"  loi-qua-10mb.xlsx: {size_mb:.1f} MB (không commit)")


# ---------------------------------------------------------------- xuất SQL


def build_sql(core: list[Member], bulk: list[Member]) -> None:
    SQL.mkdir(parents=True, exist_ok=True)
    head = ("-- Sinh tự động bởi tests/fixtures/generate.py - KHÔNG sửa tay.\n"
            "-- Tên bảng/cột lấy từ biến DB trong generate.py; nếu migration đổi tên,\n"
            "-- sửa ở đó rồi chạy lại script.\n\n")

    (SQL / "01-reset.sql").write_text(
        head + "TRUNCATE TABLE {m}, {p} RESTART IDENTITY CASCADE;\n"
                "DELETE FROM {s};\n".format(
                    m=DB["member_table"], p=DB["period_table"], s=DB["setting_table"]),
        encoding="utf-8")

    def members_sql(members: list[Member], prefix: str) -> str:
        cols = ", ".join(DB["member_cols"])
        vals = []
        for i, m in enumerate(members, start=1):
            vals.append("  ({}, {}, {}, {}, {}, '{}')".format(
                sql_str(guid(prefix, i)), sql_str(m.full_name),
                sql_date(m.date_of_birth), sql_str(m.gender),
                sql_date(m.admission), SEED_CREATED_AT))
        return f"INSERT INTO {DB['member_table']} ({cols}) VALUES\n" + ",\n".join(vals) + ";\n"

    (SQL / "02-seed-members-core.sql").write_text(
        head + f"-- {len(core)} đảng viên bộ lõi, Id tất định c0000000-...-<số thứ tự>\n"
        + members_sql(core, "c"), encoding="utf-8")

    (SQL / "03-seed-members-bulk.sql").write_text(
        head + f"-- {len(bulk)} đảng viên bộ lớn, Id tất định b0000000-...-<số thứ tự>\n"
        + members_sql(bulk, "b"), encoding="utf-8")

    pcols = ", ".join(DB["period_cols"])
    pvals = []
    for i, p in enumerate(ds.PERIODS_MAIN, start=1):
        pvals.append("  ({}, {}, {}, {}, {}, {}, '{}')".format(
            sql_str(guid("d", i)), sql_str(p.name), p.from_day, p.from_month,
            p.to_day, p.to_month, SEED_CREATED_AT))
    scols = ", ".join(DB["setting_cols"])
    s = ds.SETTINGS_DEFAULT
    setting = "INSERT INTO {t} ({c}) VALUES\n  ({i}, {a}, {b}, {st}, {u}, '{ca}');\n".format(
        t=DB["setting_table"], c=scols,
        i=sql_str("50000000-0000-4000-8000-000000000001"),
        a=s["start_years"], b=s["end_years"], st=s["step_years"],
        u=sql_str(s["unit_name"]), ca=SEED_CREATED_AT)

    (SQL / "04-seed-periods-settings.sql").write_text(
        head + f"INSERT INTO {DB['period_table']} ({pcols}) VALUES\n"
        + ",\n".join(pvals) + ";\n\n" + setting, encoding="utf-8")


# ---------------------------------------------------------------- main


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--with-oversize", action="store_true",
                    help="sinh thêm file .xlsx > 10 MB (không commit)")
    args = ap.parse_args()

    for p in ds.PERIODS_MAIN + ds.PERIODS_LEAP_EDGE + ds.PERIODS_OVERLAP + ds.PERIODS_FULL_COVER:
        p.validate()

    core = ds.MEMBERS_CORE
    bulk = ds.MEMBERS_BULK
    assert len({m.code for m in core}) == len(core), "mã ca biên bị trùng"
    assert len({m.full_name for m in core}) == len(core), "họ tên bộ lõi bị trùng"

    DATA.mkdir(parents=True, exist_ok=True)
    years = [2025, 2026, 2027, 2028]

    scen_t0 = scenario(core, ds.PERIODS_MAIN, ds.SETTINGS_DEFAULT, ds.T0, years)
    self_check(core, scen_t0)

    scen_t0_step10 = scenario(core, ds.PERIODS_MAIN, ds.SETTINGS_STEP10, ds.T0, years)
    scen_t1 = scenario(core, ds.PERIODS_MAIN, ds.SETTINGS_DEFAULT, ds.T1, years)
    scen_t2 = scenario(core, ds.PERIODS_MAIN, ds.SETTINGS_DEFAULT, ds.T2, years)
    scen_leap = scenario(core, ds.PERIODS_LEAP_EDGE, ds.SETTINGS_DEFAULT, ds.T0, years)
    scen_cover = scenario(core, ds.PERIODS_FULL_COVER, ds.SETTINGS_DEFAULT, ds.T0, years)
    scen_overlap = scenario(core, ds.PERIODS_OVERLAP, ds.SETTINGS_DEFAULT, ds.T0, years)
    scen_empty = scenario(core, [], ds.SETTINGS_DEFAULT, ds.T0, years)

    # E2E-3: đổi Bước 5 -> 10 làm danh sách Đợt 7/11 đổi 6 -> 4 người
    a = scen_t0["eligibleByPeriod"]["P4"]["byYear"]["2026"]
    b = scen_t0_step10["eligibleByPeriod"]["P4"]["byYear"]["2026"]
    assert (a["total"], b["total"]) == (6, 4), (a["total"], b["total"])
    gone = {r["code"] for r in a["rows"]} - {r["code"] for r in b["rows"]}
    assert gone == {"S01", "S03"}, gone
    # ...và badge giảm 7 -> 6 vì S04 (mốc 35 trong khoảng trống) cũng mất mốc
    assert (scen_t0["badgeCurrentYear"], scen_t0_step10["badgeCurrentYear"]) == (7, 6), \
        (scen_t0["badgeCurrentYear"], scen_t0_step10["badgeCurrentYear"])

    # E2E-4: nới Đến ngày Đợt 2/9 từ 10/09 -> 30/09 thì B03 chuyển vào đợt, badge 7 -> 6
    widened = [Period("P3", "Đợt 2/9", 15, 8, 30, 9) if p.code == "P3" else p
               for p in ds.PERIODS_MAIN]
    scen_widened = scenario(core, widened, ds.SETTINGS_DEFAULT, ds.T0, years)
    assert scen_widened["badgeCurrentYear"] == 6, scen_widened["badgeCurrentYear"]
    assert any(r["code"] == "B03"
               for r in scen_widened["eligibleByPeriod"]["P3"]["byYear"]["2026"]["rows"])

    # QT8: T1 đang diễn ra, T2 sang năm sau
    assert scen_t1["upcomingPeriod"]["status"] == "Đang diễn ra", scen_t1["upcomingPeriod"]
    assert (scen_t2["upcomingPeriod"]["code"], scen_t2["upcomingPeriod"]["year"]) == ("P1", 2027)
    assert scen_empty["upcomingPeriod"] is None
    # Không có đợt nào -> mọi người có mốc rơi trong 2026 đều bị sót (20 + 7 = 27)
    assert scen_empty["missedByYear"]["2026"]["total"] == 27, \
        scen_empty["missedByYear"]["2026"]["total"]
    assert gap_ranges(ds.PERIODS_FULL_COVER, 2026) == []
    print("  kịch bản lan truyền: OK")

    all_members = core + bulk
    search_nguyen = sum(1 for m in all_members if "Nguyễn" in m.full_name)
    total = len(all_members)

    expected = {
        "meta": {
            "generatedBy": "tests/fixtures/generate.py",
            "source": "docs/2026-09-17-huyhieudang-business-design.md v1.1, mục 3 và 4",
            "warning": "Tệp sinh tự động. Sửa dataset.py/qt_reference.py rồi chạy lại generate.py.",
        },
        "fixedDates": {"T0": iso(ds.T0), "T1": iso(ds.T1), "T2": iso(ds.T2),
                       "timezone": ds.TIMEZONE},
        "counts": {
            "core": len(core), "bulk": len(bulk), "total": total,
            "coreGenderNam": sum(1 for m in core if m.gender == "Nam"),
            "coreGenderNu": sum(1 for m in core if m.gender == "Nữ"),
            "coreGenderTrong": sum(1 for m in core if m.gender is None),
            "totalGenderNam": sum(1 for m in all_members if m.gender == "Nam"),
            "totalGenderNu": sum(1 for m in all_members if m.gender == "Nữ"),
            "totalGenderTrong": sum(1 for m in all_members if m.gender is None),
            "totalDobTrong": sum(1 for m in all_members if m.date_of_birth is None),
        },
        "pagination": {
            "totalRows": total,
            "pages": {str(size): (total + size - 1) // size for size in (10, 20, 50, 100)},
            "lastPageRows": {str(size): (total % size) or size for size in (10, 20, 50, 100)},
        },
        "search": {
            "note": "Đếm trên bộ lõi + bộ lớn đã nạp cùng nhau. Phân biệt hoa/thường: xem OQ-8.",
            "Nguyễn": search_nguyen,
            "nguyễn (viết thường)": search_nguyen,
            "Đào Văn Ân": 1,
            "Nguyễn Thị Ánh": 1,
            "Không Tồn Tại": 0,
        },
        "periods": {
            "main": [{"code": p.code, "name": p.name, "fromDay": p.from_day,
                      "fromMonth": p.from_month, "toDay": p.to_day,
                      "toMonth": p.to_month, "label": p.label()}
                     for p in ds.PERIODS_MAIN],
            "leapEdge": [{"code": p.code, "name": p.name, "fromDay": p.from_day,
                          "fromMonth": p.from_month, "toDay": p.to_day,
                          "toMonth": p.to_month, "label": p.label()}
                         for p in ds.PERIODS_LEAP_EDGE],
        },
        "exportFileNames": {
            "Đợt 7/11 năm 2026": "DuDieuKien_Dot7-11_2026.xlsx",
            "Đợt 3/2 năm 2026": "DuDieuKien_Dot3-2_2026.xlsx",
            "Đợt 19/5 năm 2026": "DuDieuKien_Dot19-5_2026.xlsx",
            "Đợt 2/9 năm 2026": "DuDieuKien_Dot2-9_2026.xlsx",
            "Chưa thuộc đợt nào năm 2026": "ChuaThuocDot_2026.xlsx",
        },
        "memberList": {
            "T0_default": member_display_rows(core, ds.SETTINGS_DEFAULT, ds.T0),
            "T0_step10": member_display_rows(core, ds.SETTINGS_STEP10, ds.T0),
        },
        "scenarios": {
            "core_default_T0": scen_t0,
            "core_step10_T0": scen_t0_step10,
            "core_default_T1": scen_t1,
            "core_default_T2": scen_t2,
            "core_default_T0_leapEdgePeriod": scen_leap,
            "core_default_T0_fullCover": scen_cover,
            "core_default_T0_overlap": scen_overlap,
            "core_default_T0_noPeriod": scen_empty,
            "core_default_T0_widenedP3": scen_widened,
        },
    }

    (DATA / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")

    def dump_members(name: str, members: list[Member]) -> None:
        (DATA / name).write_text(json.dumps(
            [{"code": m.code, "fullName": m.full_name,
              "dateOfBirth": iso(m.date_of_birth), "gender": m.gender,
              "officialAdmissionDate": iso(m.admission), "intent": m.intent}
             for m in members], ensure_ascii=False, indent=2), encoding="utf-8")

    dump_members("members-core.json", core)
    dump_members("members-bulk.json", bulk)

    (DATA / "settings.json").write_text(json.dumps(
        {"default": ds.SETTINGS_DEFAULT, "step10": ds.SETTINGS_STEP10,
         "noUnitName": ds.SETTINGS_NO_UNIT}, ensure_ascii=False, indent=2), encoding="utf-8")

    (DATA / "periods.json").write_text(json.dumps(
        expected["periods"], ensure_ascii=False, indent=2), encoding="utf-8")

    (DATA / "import-rows.json").write_text(json.dumps({
        "header": ds.HEADER,
        "errorRows": [{"fullName": r[0], "dateOfBirth": r[1], "gender": r[2],
                       "officialAdmissionDate": r[3], "errorCode": r[4],
                       "expectedReason": r[5]} for r in ds.ERR_ROWS],
        "okRows": [{"fullName": r[0], "dateOfBirth": r[1], "gender": r[2],
                    "officialAdmissionDate": r[3]} for r in ds.OK_ROWS],
        "okRowToday": {"fullName": ds.OK_ROW_TODAY[0], "dateOfBirth": ds.OK_ROW_TODAY[1],
                       "gender": ds.OK_ROW_TODAY[2],
                       "officialAdmissionDate": ds.OK_ROW_TODAY[3]},
        "undecidedRows": [{"fullName": r[0], "dateOfBirth": r[1], "gender": r[2],
                           "officialAdmissionDate": r[3], "openQuestion": r[4],
                           "question": r[5]} for r in ds.NORMALIZE_ROWS],
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    manifest = build_excel(core, bulk)
    (DATA / "excel-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    build_sql(core, bulk)

    if args.with_oversize:
        build_oversize()

    print(f"  bộ lõi {len(core)} người · bộ lớn {len(bulk)} người · tổng {total}")
    print(f"  {len(manifest)} tệp Excel · expected.json "
          f"{(DATA / 'expected.json').stat().st_size // 1024} KB")
    print("Xong.")


if __name__ == "__main__":
    os.chdir(ROOT)
    main()
